from PyPDF2 import PdfReader
import docx, pandas as pd, pptx

def extract_text_from_pdf(file):
    reader = PdfReader(file)
    return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])

def extract_text_from_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_from_excel(file):
    df = pd.read_excel(file)
    return df.to_string()

def extract_text_from_ppt(file):
    prs = pptx.Presentation(file)
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_runs.append(shape.text)
    return "\n".join(text_runs)

def extract_text_from_file(file, filename: str):
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file)
    elif filename.endswith(".xlsx"):
        return extract_text_from_excel(file)
    elif filename.endswith(".pptx"):
        return extract_text_from_ppt(file)
    else:
        return "Unsupported file type."
